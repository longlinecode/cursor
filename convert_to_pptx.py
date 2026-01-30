#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt
import re

def parse_markdown_slides(md_content):
    """Parse markdown content into slides."""
    slides = []
    current_slide = {"title": "", "subtitle": "", "content": []}

    lines = md_content.strip().split('\n')
    i = 0

    while i < len(lines):
        line = lines[i]

        # Slide separator
        if line.strip() == '---':
            if current_slide["title"] or current_slide["content"]:
                slides.append(current_slide)
                current_slide = {"title": "", "subtitle": "", "content": []}
            i += 1
            continue

        # Main title (# )
        if line.startswith('# ') and not line.startswith('## '):
            current_slide["title"] = line[2:].strip()
            i += 1
            continue

        # Subtitle (## )
        if line.startswith('## '):
            if not current_slide["subtitle"]:
                current_slide["subtitle"] = line[3:].strip()
            else:
                current_slide["content"].append(("heading", line[3:].strip()))
            i += 1
            continue

        # Bullet points
        if line.startswith('- '):
            current_slide["content"].append(("bullet", line[2:].strip()))
            i += 1
            continue

        # Numbered items
        if re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line)
            current_slide["content"].append(("number", text.strip()))
            i += 1
            continue

        # Bold text as sub-bullet
        if line.startswith('**') and '**:' in line:
            current_slide["content"].append(("bullet", line.strip()))
            i += 1
            continue

        # Regular text
        if line.strip():
            current_slide["content"].append(("text", line.strip()))

        i += 1

    # Don't forget the last slide
    if current_slide["title"] or current_slide["content"]:
        slides.append(current_slide)

    return slides

def create_pptx(slides, output_path):
    """Create PowerPoint from parsed slides."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for slide_data in slides:
        # Choose layout based on content
        if not slide_data["content"] and slide_data["subtitle"]:
            # Title slide
            slide_layout = prs.slide_layouts[0]  # Title Slide
            slide = prs.slides.add_slide(slide_layout)

            title = slide.shapes.title
            subtitle = slide.placeholders[1]

            title.text = slide_data["title"]
            subtitle.text = slide_data["subtitle"]
        else:
            # Content slide
            slide_layout = prs.slide_layouts[1]  # Title and Content
            slide = prs.slides.add_slide(slide_layout)

            # Set title
            title = slide.shapes.title
            title.text = slide_data["title"]

            # Add content
            if len(slide.placeholders) > 1:
                body = slide.placeholders[1]
                tf = body.text_frame
                tf.clear()

                first_item = True
                for content_type, text in slide_data["content"]:
                    if first_item:
                        p = tf.paragraphs[0]
                        first_item = False
                    else:
                        p = tf.add_paragraph()

                    p.text = text
                    p.level = 0

                    if content_type == "heading":
                        p.font.bold = True
                        p.font.size = Pt(24)
                    elif content_type == "bullet" or content_type == "number":
                        p.font.size = Pt(18)
                        # Handle bold text within bullets
                        if text.startswith('**'):
                            p.font.bold = True
                    else:
                        p.font.size = Pt(18)

    prs.save(output_path)
    print(f"Saved: {output_path}")

def main():
    # Read markdown file
    with open('/home/user/cursor/us-china-relations-slides.md', 'r') as f:
        md_content = f.read()

    # Parse and create PPTX
    slides = parse_markdown_slides(md_content)
    create_pptx(slides, '/home/user/cursor/us-china-relations-slides.pptx')

if __name__ == "__main__":
    main()
